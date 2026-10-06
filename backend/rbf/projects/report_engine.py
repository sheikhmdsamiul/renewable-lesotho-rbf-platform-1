"""Report engine: one registry, one dataset per report, three outputs.

Each report is registered once with its audience, category and the filters it understands.
Its builder returns a `ReportData` (summary figures plus tables). The PDF, the Excel workbook
and the CSV are all rendered from that same object, so their figures always agree.

Access is enforced here, not in the browser: a role can only generate the reports registered
for it, and every query runs inside that role's data scope (vendor: own projects; DoE: its
districts; field verifier: its own work and area). Reports that are not built yet are simply
not registered.
"""
from __future__ import annotations

import csv
import html
import io
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from typing import Callable

from django.db.models import Avg, Count, Max, Q, Sum
from django.utils import timezone
from django.utils.dateparse import parse_date
from rest_framework.exceptions import PermissionDenied, ValidationError

from rbf.users.models import UserRole

from . import claim_status
from .district_scope import doe_districts, doe_region_filter, field_verifier_district_filter, vendor_query_filter
from .models import (
    AnomalyFlag,
    AuditLog,
    FieldVerification,
    FieldVerificationStatus,
    InstallationReport,
    InstallationStatus,
    PaymentClaim,
    Project,
    ProjectStatus,
    ProspectSyncLog,
    VerificationTask,
)

R = UserRole


def inclusion_targets() -> dict:
    """Programme minimum inclusion targets (%), from Platform Configuration."""
    from rbf.users.models import inclusion_targets as configured

    return configured()


def project_inclusion_targets(project) -> dict:
    """A project's contracted inclusion targets, falling back to the programme minimums."""
    minimums = inclusion_targets()
    return {
        'female': int(project.female_target_pct or project.target_female_pct or minimums['female']),
        'vulnerable': int(project.vulnerable_target_pct or project.target_vulnerable_pct or minimums['vulnerable']),
        'low_income': int(project.low_income_target_pct or project.target_low_income_pct or minimums['low_income']),
    }
PDF_TABLE_ROWS = 60


# ---------------------------------------------------------------------------------------
# Report data
# ---------------------------------------------------------------------------------------

# Filters a report can declare. "period" is the from/to dates; the others narrow the projects
# in scope. A report only receives the filters it declares; anything else is ignored.
PROJECT_FILTERS = ('period', 'district', 'technology', 'vendor', 'tender')
FILTER_LABELS = {'district': 'District', 'technology': 'Technology', 'vendor': 'Vendor', 'tender': 'Tender', 'project': 'Project'}


@dataclass
class Table:
    title: str
    columns: list[str]
    rows: list[list]
    note: str = ''
    # (label column, value column): drawn as a bar chart in the PDF.
    chart: tuple | None = None
    # Column name -> 'identity' (beneficiary identifiers) or 'gps' (exact coordinates);
    # masked according to the viewer's role (see pii.py).
    pii: dict = field(default_factory=dict)


@dataclass
class ReportData:
    title: str
    summary: list[tuple[str, object]] = field(default_factory=list)
    details: list[tuple[str, object]] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    # Anything a format-specific renderer needs beyond the tables (for example IATI fields).
    extra: dict = field(default_factory=dict)

    @property
    def row_count(self) -> int:
        return sum(len(t.rows) for t in self.tables)


@dataclass
class ReportDefinition:
    id: str
    title: str
    description: str
    category: str
    roles: frozenset
    builder: Callable
    quick: bool = True
    requires_project: bool = False
    formats: tuple = ('pdf', 'excel', 'csv')
    filters: tuple = PROJECT_FILTERS
    method: str = ''
    # Formal reports are reviewed and approved before they are distributed (see report_workflow).
    formal: bool = False
    # Per-tender reports need the tender filter; personal-data exports need a stated reason.
    requires_tender: bool = False
    requires_reason: bool = False
    # Extra output formats: format -> renderer(data, ctx, header, notes) -> (bytes, content type, extension).
    renderers: dict = field(default_factory=dict)


class ReportContext:
    """The requesting user, their filters, and their data scope."""

    def __init__(self, user, filters: dict | None, allowed: tuple = PROJECT_FILTERS):
        self.user = user
        self.filters = filters or {}
        uses_period = 'period' in allowed
        self.date_from = self._date('from', 'date_from') if uses_period else None
        self.date_to = self._date('to', 'date_to') if uses_period else None
        if self.date_from and self.date_to and self.date_from > self.date_to:
            raise ValidationError({'filters': 'The start date is after the end date.'})
        self.project_id = str(self.filters.get('project_id') or self.filters.get('project') or '').strip() or None
        self.district = self._text('district') if 'district' in allowed else ''
        self.technology = self._text('technology') if 'technology' in allowed else ''
        self.vendor = self._text('vendor') if 'vendor' in allowed else ''
        self.tender = self._text('tender') if 'tender' in allowed else ''
        self.reason = str(self.filters.get('reason') or '').strip()[:500] if 'reason' in allowed else ''

    def _text(self, key) -> str:
        return str(self.filters.get(key) or '').strip()[:128]

    def _date(self, *keys):
        for key in keys:
            raw = str(self.filters.get(key) or '').strip()
            if raw:
                parsed = parse_date(raw)
                if parsed is None:
                    raise ValidationError({'filters': f'"{raw}" is not a valid date (use YYYY-MM-DD).'})
                return parsed
        return None

    @property
    def role(self):
        return getattr(self.user, 'role', None)

    @property
    def period_label(self) -> str:
        if not (self.date_from or self.date_to):
            return 'Programme to date'
        start = self.date_from.strftime('%d %b %Y') if self.date_from else 'Start'
        end = self.date_to.strftime('%d %b %Y') if self.date_to else 'today'
        return f'{start} to {end}'

    def filters_label(self) -> str:
        parts = []
        if self.district:
            parts.append(f'District: {self.district}')
        if self.technology:
            parts.append(f'Technology: {self.technology}')
        if self.vendor:
            name = Project.objects.filter(vendor_id=self.vendor).values_list('vendor_name', flat=True).first()
            parts.append(f'Vendor: {name or self.vendor}')
        if self.tender:
            from rbf.tenders.models import Tender

            ref = Tender.objects.filter(id=self.tender).values_list('reference_number', flat=True).first() if self.tender.isdigit() else None
            parts.append(f'Tender: {ref or self.tender}')
        if self.project_id:
            ref = Project.objects.filter(id=self.project_id).values_list('project_reference', flat=True).first() if self.project_id.isdigit() else None
            parts.append(f'Project: {ref or self.project_id}')
        return '; '.join(parts) or 'None'

    def in_period(self, queryset, field_name: str):
        tz = timezone.get_current_timezone()
        if self.date_from:
            queryset = queryset.filter(**{f'{field_name}__gte': timezone.make_aware(datetime.combine(self.date_from, time.min), tz)})
        if self.date_to:
            end = timezone.make_aware(datetime.combine(self.date_to + timedelta(days=1), time.min), tz)
            queryset = queryset.filter(**{f'{field_name}__lt': end})
        return queryset

    def scoped_projects(self):
        """Every project the role may report on, before any user-chosen filter."""
        qs = Project.objects.all().order_by('id')
        if self.role == R.VENDOR:
            qs = qs.filter(vendor_query_filter(self.user))
        elif self.role == R.DOE_OFFICER:
            qs = qs.filter(doe_region_filter(self.user))
        elif self.role == R.FIELD_VERIFIER:
            qs = qs.filter(field_verifier_district_filter(self.user, project_prefix=''))
        elif self.role == R.EVALUATION_COMMITTEE:
            qs = qs.none()
        return qs.distinct()

    def scoped_tenders(self):
        """Tenders the role may report on: programme-wide for oversight roles, the tenders of
        its own projects for vendors and DoE, and its committee assignments for evaluators."""
        from rbf.tenders.models import Tender

        qs = Tender.objects.all()
        if self.role == R.EVALUATION_COMMITTEE:
            qs = qs.filter(evaluation_committee_members__member=self.user)
        elif self.role not in {R.RBF_OFFICIAL, R.ADMIN, R.UNDP_DONOR, R.AUDITOR, R.TAC}:
            qs = qs.filter(id__in=[t for t in self.scoped_projects().values_list('tender_id', flat=True) if t])
        return qs.distinct().order_by('-created_at')

    def tenders(self):
        qs = self.scoped_tenders()
        if self.tender:
            qs = qs.filter(id=self.tender) if self.tender.isdigit() else qs.none()
        return qs

    def tender_obj(self):
        tender = self.tenders().first() if self.tender else None
        if tender is None:
            raise ValidationError({'filters': 'Choose a tender you can report on.'})
        return tender

    def projects(self):
        qs = self.scoped_projects()
        if self.project_id:
            qs = qs.filter(id=self.project_id) if self.project_id.isdigit() else qs.none()
        if self.district:
            qs = qs.filter(Q(district__iexact=self.district) | Q(region__iexact=self.district))
        if self.technology:
            qs = qs.filter(tech_type__iexact=self.technology)
        if self.vendor:
            qs = qs.filter(vendor_id=self.vendor)
        if self.tender:
            qs = qs.filter(tender_id=self.tender) if self.tender.isdigit() else qs.none()
        return qs

    def project(self):
        if not self.project_id:
            return None
        project = self.projects().first()
        if project is None:
            raise PermissionDenied('That project is not in your reporting scope.')
        return project

    def scope_label(self) -> str:
        if self.role == R.DOE_OFFICER:
            return 'Districts: ' + (', '.join(doe_districts(self.user)) or 'none assigned')
        if self.role == R.VENDOR:
            return 'Your projects'
        if self.role == R.FIELD_VERIFIER:
            return 'Your verification work'
        return 'National portfolio'


def filter_options(user) -> dict:
    """Values the viewer can filter by, drawn from the projects in their scope."""
    projects = ReportContext(user, {}).scoped_projects()
    districts = set()
    for district, region in projects.values_list('district', 'region'):
        for name in (district, region):
            if name and name.strip():
                districts.add(name.strip())
    vendors = {}
    for vendor_id, name in projects.values_list('vendor_id', 'vendor_name'):
        if vendor_id:
            vendors.setdefault(vendor_id, name or vendor_id)
    tenders = ReportContext(user, {}).scoped_tenders()
    return {
        'districts': sorted(districts, key=str.lower),
        'technologies': sorted({t for t in projects.values_list('tech_type', flat=True) if t}, key=str.lower),
        'vendors': [{'id': vid, 'name': name} for vid, name in sorted(vendors.items(), key=lambda item: str(item[1]).lower())],
        'tenders': [
            {'id': str(t.id), 'reference': t.reference_number, 'name': t.name}
            for t in tenders
        ],
        'projects': [
            {'id': str(p.id), 'reference': p.project_reference or f'Project {p.id}', 'vendor': p.vendor_name}
            for p in projects.order_by('project_reference', 'id')
        ],
    }


def _pct(part, whole) -> float:
    return round(part / whole * 100.0, 1) if whole else 0.0


def _money(value) -> str:
    return f'LSL {float(value or 0):,.2f}'


def _d(value) -> str:
    if not value:
        return ''
    if isinstance(value, datetime):
        return timezone.localtime(value).strftime('%Y-%m-%d %H:%M')
    return value.strftime('%Y-%m-%d') if hasattr(value, 'strftime') else str(value)


def _user_name(user) -> str:
    return (user.full_name or user.username) if user else 'System'


# ---------------------------------------------------------------------------------------
# Builders
# ---------------------------------------------------------------------------------------

def build_kpi_project(ctx: ReportContext) -> ReportData:
    from .kpi import KpiService

    project = ctx.project()
    if project is None:
        raise ValidationError({'filters': 'Choose the project for this report.'})
    summary = KpiService.for_project(str(project.id)).getFullKpiSummary()
    ip = summary.get('installation_progress') or {}
    gender = summary.get('gender_kpi') or {}
    energy = summary.get('energy_kpi') or {}
    uptime = summary.get('uptime_kpi') or {}
    female = float((gender.get('female_headed') or {}).get('percentage') or 0)
    vulnerable = float((gender.get('vulnerable') or {}).get('percentage') or 0)
    low_income = float((gender.get('low_income') or {}).get('percentage') or 0)
    targets = project_inclusion_targets(project)
    data = ReportData(title=f'KPI Report: {project.project_reference or project.id}')
    data.details = [
        ('Project', project.project_reference or project.id), ('Vendor', project.vendor_name),
        ('Technology', project.tech_type), ('District', project.district or project.region),
        ('Status', project.get_status_display()), ('KPIs calculated', summary.get('generated_at') or ''),
    ]
    data.summary = [
        ('Target installations', int(ip.get('target') or project.target_installations or 0)),
        ('Verified', int(ip.get('verified') or 0)),
        ('Progress', f"{float(ip.get('progress_pct') or 0):.1f}%"),
        ('Expected by now', f"{float(ip.get('expected_progress_pct') or 0):.1f}%"),
        ('Average uptime', f"{float(uptime.get('average_uptime_pct') or 0):.1f}%"),
        ('Energy vs target (this month)', f"{float(energy.get('current_month_pct') or 0):.1f}%"),
    ]
    data.tables.append(Table('Installation progress', ['Measure', 'Value'], [
        ['Submitted', int(ip.get('submitted') or 0)], ['Verified', int(ip.get('verified') or 0)],
        ['Pending verification', int(ip.get('pending') or 0)], ['Flagged', int(ip.get('flagged') or 0)],
    ]))
    data.tables.append(Table('Gender and inclusion', ['Indicator', 'Target', 'Actual', 'Result'], [
        ['Female-headed households', f"{targets['female']}%", f'{female:.1f}%', 'Met' if female >= targets['female'] else 'Below target'],
        ['Vulnerable households', f"{targets['vulnerable']}%", f'{vulnerable:.1f}%', 'Met' if vulnerable >= targets['vulnerable'] else 'Below target'],
        ['Low-income households', f"{targets['low_income']}%", f'{low_income:.1f}%', 'Met' if low_income >= targets['low_income'] else 'Below target'],
    ]))
    return data


def build_verification_summary(ctx: ReportContext) -> ReportData:
    projects = list(ctx.projects())
    visits = ctx.in_period(FieldVerification.objects.filter(installation__project__in=projects), 'verified_at')
    per_project = {
        row['installation__project_id']: row
        for row in visits.values('installation__project_id').annotate(
            visits=Count('id'),
            verified=Count('id', filter=Q(verification_status=FieldVerificationStatus.VERIFIED)),
            flagged=Count('id', filter=Q(verification_status=FieldVerificationStatus.FLAGGED)),
            partial=Count('id', filter=Q(verification_status=FieldVerificationStatus.PARTIAL)),
            mismatches=Count('id', filter=Q(location_match=False)),
            avg_distance=Avg('location_distance_meters'),
        )
    }
    reported = dict(
        InstallationReport.objects.filter(project__in=projects).values('project_id').annotate(n=Count('id')).values_list('project_id', 'n')
    )
    reverified = dict(
        VerificationTask.objects.filter(report__project__in=projects, verification_round__gt=1)
        .values('report__project_id').annotate(n=Count('id')).values_list('report__project_id', 'n')
    )
    rows = []
    for project in projects:
        row = per_project.get(project.id, {})
        if not row and not reported.get(project.id):
            continue
        rows.append([
            project.project_reference or project.id, project.vendor_name, project.district or project.region,
            reported.get(project.id, 0), row.get('visits', 0), row.get('verified', 0), row.get('flagged', 0),
            row.get('partial', 0), reverified.get(project.id, 0), row.get('mismatches', 0),
            round(float(row.get('avg_distance') or 0), 1),
        ])
    total_visits = visits.count()
    verified = visits.filter(verification_status=FieldVerificationStatus.VERIFIED).count()
    flagged = visits.filter(verification_status=FieldVerificationStatus.FLAGGED).count()
    data = ReportData(title='Field Verification Report')
    data.summary = [
        ('Installations reported', sum(reported.values())), ('Verification visits', total_visits),
        ('Verified', verified), ('Flagged', flagged),
        ('Pass rate', f'{_pct(verified, total_visits):.1f}%'), ('Re-verified installations', sum(reverified.values())),
    ]
    data.tables.append(Table(
        'By project',
        ['Project', 'Vendor', 'District', 'Installations reported', 'Visits', 'Verified', 'Flagged', 'Partial', 'Re-verified', 'GPS mismatches', 'Avg GPS distance (m)'],
        rows, chart=(0, 5),
    ))
    flagged_rows = [
        [fv.installation.project.project_reference or fv.installation.project_id, fv.installation.serial_number,
         fv.installation.beneficiary_id, float(fv.installation.gps_lat), float(fv.installation.gps_lng),
         fv.get_verification_status_display(), fv.flag_reason or '', round(float(fv.location_distance_meters or 0), 1),
         fv.verification_round, _user_name(fv.field_officer), _d(fv.verified_at)]
        for fv in visits.exclude(verification_status=FieldVerificationStatus.VERIFIED)
        .select_related('installation__project', 'field_officer').order_by('-verified_at')
    ]
    data.tables.append(Table(
        'Flagged and partial verifications',
        ['Project', 'Installation serial', 'Beneficiary ID', 'Latitude', 'Longitude', 'Outcome', 'Reason', 'GPS distance (m)', 'Round', 'Verifier', 'Date'],
        flagged_rows,
        pii={'Beneficiary ID': 'identity', 'Latitude': 'gps', 'Longitude': 'gps'},
    ))
    return data


def _scoped_claims(ctx: ReportContext):
    claims = PaymentClaim.objects.filter(project__in=ctx.projects()).select_related('project', 'vendor', 'milestone')
    return ctx.in_period(claims, 'submitted_at').order_by('-submitted_at')


def _claim_totals(claims) -> dict:
    totals = {'claimed': 0.0, 'paid': 0.0, 'awaiting': 0.0, 'review': 0.0, 'held': 0.0, 'rejected': 0.0}
    for claim in claims:
        amount = float(claim.claim_amount or 0)
        totals['claimed'] += amount
        if claim.status in claim_status.PAID:
            totals['paid'] += amount
        elif claim.status in claim_status.APPROVED_AWAITING_PAYMENT:
            totals['awaiting'] += amount
        elif claim.status in claim_status.HELD:
            totals['held'] += amount
        elif claim.status in claim_status.REJECTED:
            totals['rejected'] += amount
        else:
            totals['review'] += amount
    return totals


def _claim_rows(claims) -> list[list]:
    return [
        [
            claim.id, claim.project.project_reference or claim.project_id, claim.project.vendor_name,
            f'M{claim.milestone.milestone_number}' if claim.milestone_id else '', float(claim.claim_amount or 0),
            claim.status, claim_status.stage(claim.status), _d(claim.submitted_at), _d(claim.verified_at),
            _d(claim.approved_at), _d(claim.paid_at), claim.payment_reference,
        ]
        for claim in claims
    ]


CLAIM_COLUMNS = [
    'Claim', 'Project', 'Vendor', 'Milestone', 'Amount (LSL)', 'Status', 'Stage', 'Submitted', 'RMT approved',
    'TAC endorsed', 'Paid', 'Payment reference',
]


def build_financial_disbursement(ctx: ReportContext) -> ReportData:
    claims = list(_scoped_claims(ctx))
    totals = _claim_totals(claims)
    contracted = float(ctx.projects().aggregate(total=Sum('budget'))['total'] or 0)
    data = ReportData(title='Financial Disbursement Report')
    data.summary = [
        ('Contracted (project budgets)', _money(contracted)), ('Claimed in period', _money(totals['claimed'])),
        ('Paid', _money(totals['paid'])), ('Approved, awaiting payment', _money(totals['awaiting'])),
        ('In review', _money(totals['review'])), ('Held for audit', _money(totals['held'])),
    ]
    by_project: dict = {}
    for claim in claims:
        entry = by_project.setdefault(claim.project_id, [claim.project.project_reference or claim.project_id, claim.project.vendor_name, float(claim.project.budget or 0), 0.0, 0.0, 0.0])
        entry[3] += float(claim.claim_amount or 0)
        if claim.status in claim_status.PAID:
            entry[4] += float(claim.claim_amount or 0)
        elif claim.status in claim_status.APPROVED_AWAITING_PAYMENT:
            entry[5] += float(claim.claim_amount or 0)
    data.tables.append(Table('By project', ['Project', 'Vendor', 'Budget (LSL)', 'Claimed (LSL)', 'Paid (LSL)', 'Awaiting payment (LSL)'], list(by_project.values()), chart=(0, 4)))
    data.tables.append(Table('Payment claims', CLAIM_COLUMNS, _claim_rows(claims)))
    return data


def build_vendor_payment_trail(ctx: ReportContext) -> ReportData:
    claims = list(_scoped_claims(ctx))
    vendors: dict = {}
    for claim in claims:
        name = claim.project.vendor_name or claim.vendor.organization_name or claim.vendor.username
        entry = vendors.setdefault(name, [name, 0, 0, 0, 0.0, 0.0])
        entry[1] += 1
        if claim.status in claim_status.APPROVED:
            entry[2] += 1
        if claim.status in claim_status.PAID:
            entry[3] += 1
            entry[5] += float(claim.claim_amount or 0)
        entry[4] += float(claim.claim_amount or 0)
    data = ReportData(title='Vendor Payment Trail')
    totals = _claim_totals(claims)
    data.summary = [('Vendors', len(vendors)), ('Claims', len(claims)), ('Claimed', _money(totals['claimed'])), ('Paid', _money(totals['paid']))]
    data.tables.append(Table('By vendor', ['Vendor', 'Claims', 'Approved by PSC', 'Paid', 'Claimed (LSL)', 'Paid (LSL)'],
                             sorted(vendors.values(), key=lambda r: -r[4]), chart=(0, 5)))
    data.tables.append(Table('Payment trail', CLAIM_COLUMNS, _claim_rows(claims)))
    return data


def build_portfolio_summary(ctx: ReportContext) -> ReportData:
    projects = list(ctx.projects().order_by('-progress', 'id'))
    completed = sum(1 for p in projects if p.status in {ProjectStatus.COMPLETED, ProjectStatus.LEGACY_COMPLETED})
    active = sum(1 for p in projects if p.status == ProjectStatus.ACTIVE)
    data = ReportData(title='Portfolio Summary')
    data.summary = [
        ('Projects', len(projects)), ('Active', active), ('Completed', completed),
        ('Budget', _money(sum(float(p.budget or 0) for p in projects))),
        ('Average progress', f"{(sum(p.progress or 0 for p in projects) / len(projects)) if projects else 0:.1f}%"),
    ]
    data.tables.append(Table('Projects', ['Project', 'Vendor', 'Technology', 'District', 'Status', 'Progress %', 'Female-headed %', 'Uptime %', 'Budget (LSL)', 'Archived'], [
        [p.project_reference or p.id, p.vendor_name, p.tech_type, p.district or p.region, p.get_status_display(),
         p.progress or 0, round(float(p.gender_impact or 0), 1), round(float(p.uptime or 0), 1), float(p.budget or 0),
         'Yes' if p.archived_at else 'No']
        for p in projects
    ]))
    return data


def build_anomaly_report(ctx: ReportContext) -> ReportData:
    flags = ctx.in_period(AnomalyFlag.objects.filter(project__in=ctx.projects()).select_related('project', 'installation'), 'created_at')
    total = flags.count()
    open_count = flags.filter(is_resolved=False).count()
    data = ReportData(title='Anomaly Flags Report')
    data.summary = [('Flags raised', total), ('Open', open_count), ('Resolved', total - open_count)]
    data.tables.append(Table('By type', ['Type', 'Raised', 'Open'], chart=(0, 1), rows=[
        [row['flag_type'], row['n'], row['open']]
        for row in flags.values('flag_type').annotate(n=Count('id'), open=Count('id', filter=Q(is_resolved=False))).order_by('-n')
    ]))
    data.tables.append(Table('Flags', ['Flag', 'Project', 'Installation serial', 'Type', 'Status', 'Raised', 'Resolved', 'Description'], [
        [f.id, f.project.project_reference or f.project_id if f.project_id else '', f.installation.serial_number if f.installation_id else '',
         f.flag_type, 'Resolved' if f.is_resolved else 'Open', _d(f.created_at), _d(f.resolved_at), f.description or '']
        for f in flags.order_by('-created_at')
    ]))
    return data


def _inclusion_row(label, queryset):
    total = queryset.count()
    female = queryset.filter(household_type__icontains='female').count()
    vulnerable = queryset.filter(household_type__icontains='vulnerable').count()
    low = queryset.filter(household_type__icontains='low').count()
    return [label, total, _pct(female, total), _pct(vulnerable, total), _pct(low, total)]


def build_gender_impact(ctx: ReportContext) -> ReportData:
    verified = ctx.in_period(
        InstallationReport.objects.filter(project__in=ctx.projects(), status=InstallationStatus.VERIFIED), 'submitted_at',
    )
    overall = _inclusion_row('All verified installations', verified)
    targets = inclusion_targets()
    data = ReportData(title='Gender and Inclusion Report')
    data.summary = [
        ('Verified installations', overall[1]),
        ('Female-headed', f'{overall[2]:.1f}% (target {targets["female"]}%)'),
        ('Vulnerable', f'{overall[3]:.1f}% (target {targets["vulnerable"]}%)'),
        ('Low-income', f'{overall[4]:.1f}% (target {targets["low_income"]}%)'),
    ]
    cols = ['Group', 'Verified installations', 'Female-headed %', 'Vulnerable %', 'Low-income %']
    techs = sorted({t for t in verified.values_list('project__tech_type', flat=True) if t})
    districts = sorted({d for d in verified.values_list('district', flat=True) if d})
    vendors = sorted({v for v in verified.values_list('project__vendor_name', flat=True) if v})
    data.tables.append(Table('By technology', cols, [_inclusion_row(t, verified.filter(project__tech_type=t)) for t in techs]))
    data.tables.append(Table('By district', cols, [_inclusion_row(d, verified.filter(district=d)) for d in districts], chart=(0, 2)))
    data.tables.append(Table('By vendor', cols, [_inclusion_row(v, verified.filter(project__vendor_name=v)) for v in vendors]))
    return data


def build_regional_progress(ctx: ReportContext) -> ReportData:
    projects = ctx.projects()
    counts = {
        row['project_id']: row
        for row in InstallationReport.objects.filter(project__in=projects).values('project_id').annotate(
            total=Count('id'),
            verified=Count('id', filter=Q(status=InstallationStatus.VERIFIED)),
            flagged=Count('id', filter=Q(status=InstallationStatus.FLAGGED)),
        )
    }
    rows = []
    districts: dict = {}
    for p in projects.order_by('district', 'id'):
        c = counts.get(p.id, {})
        district = p.district or p.region or 'Unknown'
        rows.append([p.project_reference or p.id, p.vendor_name, district, p.tech_type, p.get_status_display(), p.progress or 0,
                     p.target_installations or 0, c.get('total', 0), c.get('verified', 0), c.get('flagged', 0)])
        d = districts.setdefault(district, [district, 0, 0, 0, 0])
        d[1] += 1
        d[2] += c.get('total', 0)
        d[3] += c.get('verified', 0)
        d[4] += c.get('flagged', 0)
    data = ReportData(title='Regional Progress Report')
    data.details = [('Scope', ctx.scope_label())]
    data.summary = [
        ('Projects', len(rows)), ('Installations reported', sum(r[7] for r in rows)),
        ('Verified', sum(r[8] for r in rows)), ('Flagged', sum(r[9] for r in rows)),
    ]
    data.tables.append(Table('By district', ['District', 'Projects', 'Installations', 'Verified', 'Flagged'], sorted(districts.values()), chart=(0, 3)))
    data.tables.append(Table('Projects', ['Project', 'Vendor', 'District', 'Technology', 'Status', 'Progress %', 'Target', 'Installations', 'Verified', 'Flagged'], rows))
    return data


def build_regional_kpi(ctx: ReportContext) -> ReportData:
    projects = ctx.projects()
    installs = ctx.in_period(InstallationReport.objects.filter(project__in=projects), 'submitted_at')
    verified = installs.filter(status=InstallationStatus.VERIFIED)
    inclusion = _inclusion_row('Region', verified)
    stats = projects.aggregate(uptime=Avg('uptime'), progress=Avg('progress'), energy=Sum('energy_output'))
    data = ReportData(title='Regional KPI Report')
    data.details = [('Scope', ctx.scope_label())]
    data.summary = [
        ('Installations reported', installs.count()), ('Verified', verified.count()),
        ('Female-headed', f'{inclusion[2]:.1f}%'), ('Vulnerable', f'{inclusion[3]:.1f}%'), ('Low-income', f'{inclusion[4]:.1f}%'),
        ('Average project uptime', f"{float(stats['uptime'] or 0):.1f}%"),
    ]
    data.tables.append(Table('By district', ['District', 'Verified installations', 'Female-headed %', 'Vulnerable %', 'Low-income %'], [
        _inclusion_row(d, verified.filter(district=d)) for d in sorted({d for d in verified.values_list('district', flat=True) if d})
    ]))
    data.tables.append(Table('By project', ['Project', 'Technology', 'Progress %', 'Uptime %', 'Energy output (kWh)', 'Female-headed %'], [
        [p.project_reference or p.id, p.tech_type, p.progress or 0, round(float(p.uptime or 0), 1), round(float(p.energy_output or 0), 1), round(float(p.gender_impact or 0), 1)]
        for p in projects.order_by('id')
    ]))
    return data


def build_technology_breakdown(ctx: ReportContext) -> ReportData:
    projects = ctx.projects()
    installs = dict(
        InstallationReport.objects.filter(project__in=projects).values('project__tech_type').annotate(n=Count('id')).values_list('project__tech_type', 'n')
    )
    verified = dict(
        InstallationReport.objects.filter(project__in=projects, status=InstallationStatus.VERIFIED)
        .values('project__tech_type').annotate(n=Count('id')).values_list('project__tech_type', 'n')
    )
    rows = [
        [row['tech_type'] or 'Unknown', row['total'], row['active'], row['completed'], round(float(row['avg_progress'] or 0), 1),
         installs.get(row['tech_type'], 0), verified.get(row['tech_type'], 0), float(row['budget'] or 0)]
        for row in projects.values('tech_type').annotate(
            total=Count('id'), active=Count('id', filter=Q(status=ProjectStatus.ACTIVE)),
            completed=Count('id', filter=Q(status__in=[ProjectStatus.COMPLETED, ProjectStatus.LEGACY_COMPLETED])),
            avg_progress=Avg('progress'), budget=Sum('budget'),
        ).order_by('-total')
    ]
    data = ReportData(title='Technology Breakdown')
    data.details = [('Scope', ctx.scope_label())]
    data.summary = [('Technologies', len(rows)), ('Projects', sum(r[1] for r in rows)), ('Verified installations', sum(r[6] for r in rows))]
    data.tables.append(Table('By technology', ['Technology', 'Projects', 'Active', 'Completed', 'Avg progress %', 'Installations', 'Verified', 'Budget (LSL)'], rows, chart=(0, 6)))
    return data


def _own_visits(ctx: ReportContext):
    return FieldVerification.objects.filter(field_officer=ctx.user).select_related('installation__project')


def _visit_rows(visits) -> list[list]:
    return [
        [fv.installation.serial_number, fv.installation.project.project_reference or fv.installation.project_id,
         fv.installation.district or fv.installation.project.district, fv.get_verification_status_display(), fv.verification_round,
         round(float(fv.location_distance_meters or 0), 1), 'Yes' if fv.location_match else 'No', fv.flag_reason or '', _d(fv.verified_at)]
        for fv in visits
    ]


VISIT_COLUMNS = ['Installation serial', 'Project', 'District', 'Outcome', 'Round', 'GPS distance (m)', 'Location match', 'Flag reason', 'Verified at']


def build_fo_history(ctx: ReportContext) -> ReportData:
    visits = ctx.in_period(_own_visits(ctx), 'verified_at').order_by('-verified_at')
    total = visits.count()
    verified = visits.filter(verification_status=FieldVerificationStatus.VERIFIED).count()
    data = ReportData(title='My Verification History')
    data.details = [('Verifier', _user_name(ctx.user))]
    data.summary = [('Visits', total), ('Verified', verified), ('Flagged or partial', total - verified)]
    data.tables.append(Table('Visits', VISIT_COLUMNS, _visit_rows(visits)))
    return data


def build_fo_daily(ctx: ReportContext) -> ReportData:
    day = ctx.date_to or timezone.localdate()
    tz = timezone.get_current_timezone()
    start = timezone.make_aware(datetime.combine(day, time.min), tz)
    visits = _own_visits(ctx).filter(verified_at__gte=start, verified_at__lt=start + timedelta(days=1)).order_by('verified_at')
    pending = VerificationTask.objects.filter(assigned_verifier=ctx.user, status__in=['Pending', 'Reverification Required']).count()
    data = ReportData(title=f'Daily Summary: {day:%d %b %Y}')
    data.details = [('Verifier', _user_name(ctx.user)), ('Day', day.strftime('%A %d %B %Y'))]
    data.summary = [
        ('Visits completed', visits.count()),
        ('Verified', visits.filter(verification_status=FieldVerificationStatus.VERIFIED).count()),
        ('Flagged or partial', visits.exclude(verification_status=FieldVerificationStatus.VERIFIED).count()),
        ('Tasks still assigned to you', pending),
    ]
    data.tables.append(Table('Visits completed', VISIT_COLUMNS, _visit_rows(visits)))
    return data


def build_fo_performance(ctx: ReportContext) -> ReportData:
    visits = ctx.in_period(_own_visits(ctx), 'verified_at')
    total = visits.count()
    verified = visits.filter(verification_status=FieldVerificationStatus.VERIFIED).count()
    avg_gps = visits.aggregate(avg=Avg('location_distance_meters'))['avg']
    data = ReportData(title='My Performance Summary')
    data.details = [('Verifier', _user_name(ctx.user))]
    data.summary = [
        ('Visits', total), ('Verified', verified), ('Pass rate', f'{_pct(verified, total):.1f}%'),
        ('Average GPS distance', f'{float(avg_gps or 0):.1f} m'), ('Location mismatches', visits.filter(location_match=False).count()),
    ]
    monthly: dict = {}
    for fv in visits.values('verified_at', 'verification_status'):
        key = timezone.localtime(fv['verified_at']).strftime('%Y-%m')
        entry = monthly.setdefault(key, [key, 0, 0])
        entry[1] += 1
        entry[2] += fv['verification_status'] == FieldVerificationStatus.VERIFIED
    data.tables.append(Table('By month', ['Month', 'Visits', 'Verified'], sorted(monthly.values(), reverse=True), chart=(0, 1)))
    data.tables.append(Table('Visits', VISIT_COLUMNS, _visit_rows(visits.order_by('-verified_at'))))
    return data


def build_full_audit(ctx: ReportContext) -> ReportData:
    logs = ctx.in_period(AuditLog.objects.select_related('actor'), 'created_at').order_by('-created_at')
    total = logs.count()
    data = ReportData(title='Audit Trail Report')
    data.summary = [('Events', total), ('People', logs.exclude(actor__isnull=True).values('actor').distinct().count())]
    data.tables.append(Table('By module', ['Module', 'Events', 'Last event'], chart=(0, 1), rows=[
        [row['module'] or '-', row['n'], _d(row['last'])]
        for row in logs.values('module').annotate(n=Count('id'), last=Max('created_at')).order_by('-n')
    ]))
    data.tables.append(Table('Events', ['Time', 'Actor', 'Role', 'Action', 'Module', 'Record', 'From', 'To', 'Notes', 'IP address'], [
        [_d(log.created_at), _user_name(log.actor) if log.actor_id else 'System', log.actor_role, log.action, log.module,
         f'{log.entity_type}#{log.entity_id}' if log.entity_id else log.entity_type, log.old_status, log.new_status, log.notes, log.ip_address]
        for log in logs
    ]))
    return data


def build_data_integrity(ctx: ReportContext) -> ReportData:
    installs = InstallationReport.objects.all()
    duplicate_serials = installs.values('serial_number').annotate(n=Count('id')).filter(n__gt=1).count()
    duplicate_beneficiaries = installs.exclude(beneficiary_id='').values('beneficiary_id').annotate(n=Count('id')).filter(n__gt=1).count()
    checks = [
        ['Projects without a reference', Project.objects.filter(Q(project_reference='') | Q(project_reference__isnull=True)).count()],
        ['Projects without a vendor', Project.objects.filter(Q(vendor_name='') | Q(vendor_id='')).count()],
        ['Installations without GPS', installs.filter(Q(gps_lat=0) | Q(gps_lng=0)).count()],
        ['Serial numbers used more than once', duplicate_serials],
        ['Beneficiary IDs used more than once', duplicate_beneficiaries],
        ['Verified installations with no field verification', installs.filter(status=InstallationStatus.VERIFIED, field_verifications__isnull=True).count()],
        ['Paid claims with no payment reference', PaymentClaim.objects.filter(status__in=claim_status.PAID, payment_reference='').count()],
        ['Open anomaly flags', AnomalyFlag.objects.filter(is_resolved=False).count()],
        ['Prospect sync failures', ProspectSyncLog.objects.filter(status='failed').count()],
    ]
    data = ReportData(title='Data Integrity Report')
    data.summary = [
        ('Projects', Project.objects.count()), ('Installations', installs.count()),
        ('Field verifications', FieldVerification.objects.count()), ('Payment claims', PaymentClaim.objects.count()),
    ]
    data.tables.append(Table('Integrity checks', ['Check', 'Records', 'Result'], [[c, n, 'Pass' if n == 0 else 'Review'] for c, n in checks]))
    return data


def build_prospect_sync(ctx: ReportContext) -> ReportData:
    logs = ctx.in_period(ProspectSyncLog.objects.all(), 'created_at').order_by('-created_at')
    total = logs.count()
    data = ReportData(title='Prospect Sync Audit')
    data.summary = [
        ('Sync attempts', total), ('Succeeded', logs.filter(status='success').count()),
        ('Failed', logs.filter(status='failed').count()), ('Pending', logs.filter(status='pending').count()),
    ]
    data.tables.append(Table('By method', ['Method', 'Attempts', 'Failed'], [
        [row['method_name'], row['n'], row['failed']]
        for row in logs.values('method_name').annotate(n=Count('id'), failed=Count('id', filter=Q(status='failed'))).order_by('-n')
    ]))
    data.tables.append(Table('Sync log', ['Log', 'Method', 'Record', 'Status', 'Attempts', 'Created', 'Updated', 'Error'], [
        [log.id, log.method_name, f'{log.record_type}#{log.record_id}' if log.record_id else log.record_type, log.status,
         log.attempts, _d(log.created_at), _d(log.updated_at), log.error_message]
        for log in logs
    ]))
    return data


# ---------------------------------------------------------------------------------------
# Registry
# ---------------------------------------------------------------------------------------



def definitions_for(user) -> list[ReportDefinition]:
    from .report_catalogue import definitions_for as _definitions_for

    return _definitions_for(user)


def definition_for(user, report_id: str) -> ReportDefinition:
    from .report_catalogue import definition_for as _definition_for

    return _definition_for(user, report_id)


# ---------------------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------------------

def _cell(value):
    if value is None:
        return ''
    if isinstance(value, float):
        # Builders round money and percentages themselves; 6 places keeps exact GPS coordinates.
        return round(value, 6)
    return value


def _apply_pii(data: ReportData, ctx: ReportContext) -> bool:
    """Mask personal-data columns for this viewer. Returns True when anything was masked."""
    from .pii import masker_for

    masked = False
    for table in data.tables:
        for column, kind in table.pii.items():
            masker = masker_for(ctx.role, kind)
            if masker is None or column not in table.columns:
                continue
            idx = table.columns.index(column)
            for row in table.rows:
                row[idx] = masker(row[idx])
            masked = True
    return masked


def _header_lines(data: ReportData, ctx: ReportContext, definition: ReportDefinition, reference: str) -> list[tuple[str, object]]:
    return [
        ('Report', data.title),
        ('Reference', reference),
        ('Period', ctx.period_label if 'period' in definition.filters else 'Not period-based (see Methodology)'),
        ('Scope', ctx.scope_label()),
        ('Filters', ctx.filters_label()),
        *data.details,
        ('Generated by', f"{_user_name(ctx.user)} ({ctx.role})"),
        ('Data extracted', timezone.localtime(timezone.now()).strftime('%Y-%m-%d %H:%M %Z')),
    ]


def _notes(data: ReportData, definition: ReportDefinition, masked: bool) -> list[tuple[str, str]]:
    from .pii import POLICY_SUMMARY

    notes = []
    if definition.method:
        notes.append(('Methodology', definition.method))
    notes.extend(('Note', n) for n in data.notes)
    if masked:
        notes.append(('Data protection', 'Some personal data in this report is masked for your role. ' + POLICY_SUMMARY))
    return notes


def render_csv(data: ReportData, ctx: ReportContext) -> bytes:
    """The report's detail table (its last table) as plain CSV, ready for analysis."""
    table = data.tables[-1] if data.tables else Table(data.title, ['Measure', 'Value'], [[k, v] for k, v in data.summary])
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(table.columns)
    for row in table.rows:
        writer.writerow([_cell(v) for v in row])
    return buffer.getvalue().encode('utf-8-sig')


def render_xlsx(data: ReportData, ctx: ReportContext, header: list, notes: list) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    bold = Font(bold=True)
    for label, value in header:
        ws.append([label, str(value)])
        ws.cell(ws.max_row, 1).font = bold
    ws.append([])
    for label, value in data.summary:
        ws.append([label, _cell(value)])
    for label, text in notes:
        ws.append([])
        ws.append([label, text])
        ws.cell(ws.max_row, 1).font = bold
        ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 90
    used = {'Summary'}
    for table in data.tables:
        name = ''.join(ch for ch in table.title if ch not in '[]:*?/\\')[:31] or 'Data'
        base, n = name, 2
        while name in used:
            name = f'{base[:28]} {n}'
            n += 1
        used.add(name)
        sheet = wb.create_sheet(name)
        sheet.append(table.columns)
        for cell in sheet[1]:
            cell.font = bold
        for row in table.rows:
            sheet.append([_cell(v) for v in row])
        sheet.freeze_panes = 'A2'
        if table.rows:
            sheet.auto_filter.ref = sheet.dimensions
        for idx, col in enumerate(table.columns, start=1):
            sheet.column_dimensions[sheet.cell(1, idx).column_letter].width = max(12, min(40, len(str(col)) + 4))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _bar_chart_svg(table: Table, limit: int = 12) -> str:
    """A horizontal bar chart of one numeric column, largest first."""
    label_idx, value_idx = table.chart
    points = []
    for row in table.rows:
        try:
            points.append((str(row[label_idx]), float(row[value_idx] or 0)))
        except (TypeError, ValueError, IndexError):
            continue
    points = sorted(points, key=lambda p: -p[1])[:limit]
    if not points or max(v for _, v in points) <= 0:
        return ''
    peak = max(v for _, v in points)
    bar_h, gap, label_w, width = 16, 8, 170, 640
    plot_w = width - label_w - 90
    height = len(points) * (bar_h + gap) + 24
    bars = []
    for i, (label, value) in enumerate(points):
        y = 8 + i * (bar_h + gap)
        w = max(2.0, plot_w * value / peak)
        shown = f'{value:,.0f}' if value >= 100 or float(value).is_integer() else f'{value:,.1f}'
        bars.append(
            f'<text x="{label_w - 8}" y="{y + 12}" text-anchor="end" font-size="10" fill="#334155">{html.escape(label[:28])}</text>'
            f'<rect x="{label_w}" y="{y}" width="{w:.1f}" height="{bar_h}" rx="3" fill="#0f766e"></rect>'
            f'<text x="{label_w + w + 6:.1f}" y="{y + 12}" font-size="10" fill="#0f172a">{shown}</text>'
        )
    caption = html.escape(f'{table.columns[value_idx]} by {table.columns[label_idx].lower()}' + (f' (top {limit})' if len(table.rows) > limit else ''))
    return (
        f'<p class="report-meta">{caption}</p>'
        f'<svg width="100%" viewBox="0 0 {width} {height}" xmlns="http://www.w3.org/2000/svg" style="max-width:{width}px;">'
        f'<line x1="{label_w}" y1="0" x2="{label_w}" y2="{height - 16}" stroke="#cbd5e1"></line>{"".join(bars)}</svg>'
    )


WIDE_TABLE_COLUMNS = 8


def _display(value):
    """A cell as printed in the PDF: thousands separators, no trailing ".0"."""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return '' if value is None else value
    if isinstance(value, int) or float(value).is_integer():
        return f'{int(value):,}'
    if abs(value) >= 1000:
        return f'{value:,.2f}'
    return f'{value:.6f}'.rstrip('0').rstrip('.')


def render_pdf_html(data: ReportData, ctx: ReportContext, header: list, notes: list) -> str:
    from .report_templates import render_kv_list, render_report_html, render_section, render_stats_grid, render_table

    wide = any(len(t.columns) > WIDE_TABLE_COLUMNS for t in data.tables)
    # Wide tables print on landscape pages; charts and table rows are not split across pages.
    page_css = (
        '<style>'
        + ('@page { size: A4 landscape; } ' if wide else '')
        + 'svg { break-inside: avoid; page-break-inside: avoid; } tr { break-inside: avoid; } '
        + '.wide-table table { font-size: 9px; } .wide-table th, .wide-table td { padding: 5px 6px; } '
        + '</style>'
    )
    parts = [page_css, render_section('About this report', render_kv_list(header))]
    if data.summary:
        parts.append(render_section('Summary', render_stats_grid([{'label': k, 'value': _display(v)} for k, v in data.summary])))
    for table in data.tables:
        rows = [[_display(_cell(v)) for v in row] for row in table.rows]
        shown = rows[:PDF_TABLE_ROWS]
        body = _bar_chart_svg(table) if table.chart else ''
        grid = render_table(table.columns, shown)
        body += f'<div class="wide-table">{grid}</div>' if len(table.columns) > WIDE_TABLE_COLUMNS else grid
        if len(rows) > len(shown):
            body += f'<p class="report-meta">Showing the first {len(shown)} of {len(rows)} rows. The Excel and CSV versions contain every row.</p>'
        parts.append(render_section(table.title, body))
    if notes:
        parts.append(render_section('Methodology and notes', ''.join(
            f'<p><strong>{html.escape(label)}:</strong> {html.escape(text)}</p>' for label, text in notes
        )))
    meta = {'Reference': header[1][1], 'Period': header[2][1], 'Scope': ctx.scope_label()}
    return render_report_html(data.title, ''.join(parts), meta)


def build(user, report_id: str, filters: dict | None):
    """Validate access and build the report data (no rendering)."""
    definition = definition_for(user, report_id)
    ctx = ReportContext(user, filters, allowed=definition.filters)
    if definition.requires_project and not ctx.project_id:
        raise ValidationError({'filters': 'Choose the project for this report.'})
    if definition.requires_tender:
        ctx.tender_obj()
    if definition.requires_reason and len(ctx.reason) < 10:
        raise ValidationError({'filters': 'State why you need this personal data (at least 10 characters). The reason is recorded in the audit log.'})
    return definition, ctx


def generate(user, report_id: str, fmt: str, filters: dict | None, reference: str = ''):
    """Build and render a report. Returns (definition, data, content_bytes, content_type, extension)."""
    definition, ctx = build(user, report_id, filters)
    if fmt not in definition.formats:
        raise ValidationError({'format': f'This report is available as {", ".join(definition.formats)}.'})
    data = definition.builder(ctx)
    masked = _apply_pii(data, ctx)
    reference = reference or f'RPT-{timezone.localdate():%Y%m%d}'
    header = _header_lines(data, ctx, definition, reference)
    notes = _notes(data, definition, masked)
    if fmt == 'csv':
        return definition, data, render_csv(data, ctx), 'text/csv', 'csv'
    if fmt == 'excel':
        return definition, data, render_xlsx(data, ctx, header, notes), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'xlsx'
    if fmt in definition.renderers:
        content, ctype, ext = definition.renderers[fmt](data, ctx, header, notes)
        return definition, data, content, ctype, ext
    from .report_pdf import render_report_pdf

    pdf = render_report_pdf(render_pdf_html(data, ctx, header, notes), title=data.title, viewer=_user_name(user), reference=reference)
    return definition, data, pdf, 'application/pdf', 'pdf'
