"""Second- and third-wave reports: procurement, oversight, exports (GeoJSON, IATI XML, ZIP) and access logs."""

import csv
import hashlib
import io
import json
import zipfile
from datetime import timedelta
from unittest import mock
from xml.etree import ElementTree as ET

from django.utils import timezone
from rest_framework import status
from rest_framework.test import APITestCase

from rbf.tenders.models import (
    EvaluationConflictOfInterest,
    Tender,
    TenderBid,
    TenderBidEvaluation,
    TenderEvaluationCommitteeMember,
)
from rbf.users.models import User
from . import report_queue
from .models import AuditLog, FieldVerification, InstallationReport, InstallationStatus, PaymentClaim, Project


def make_user(username, role, **extra):
    return User.objects.create_user(username=username, password="securePass123", role=role, status="Active", **extra)


@mock.patch("rbf.projects.report_exports._dossier_pdf", return_value=(b"%PDF-1.4 dossier", "Rendered for the test."))
@mock.patch("rbf.projects.report_pdf.render_report_pdf", return_value=b"%PDF-1.4 test")
class SecondWaveReportTests(APITestCase):
    def setUp(self):
        self.admin = make_user("p4_admin", "Platform Administrator (Super Admin)")
        self.rmt = make_user("p4_rmt", "RBF Management Team")
        self.tac = make_user("p4_tac", "TAC Member")
        self.psc = make_user("p4_psc", "Project Steering Committee")
        self.auditor = make_user("p4_auditor", "Auditor")
        self.doe = make_user("p4_doe", "DoE Officer", region="Maseru")
        self.vendor = make_user("p4_vendor", "Vendor", organization_name="Sun Co")
        self.ec1 = make_user("p4_ec1", "Evaluation Committee")
        self.ec2 = make_user("p4_ec2", "Evaluation Committee")
        now = timezone.now()
        self.tender = Tender.objects.create(reference_number="TND-P4-1", name="Solar homes", department="DoE", category="Energy",
                                            deadline=now - timedelta(days=20), status="Awarded", published_at=now - timedelta(days=40),
                                            awarded_at=now - timedelta(days=5), cooling_off_until=now - timedelta(days=2),
                                            awarded_vendor_name="Sun Co", budget=100000, technology_types=["SHS"])
        self.other_tender = Tender.objects.create(reference_number="TND-P4-2", name="Cookstoves", department="DoE", category="Energy",
                                                  deadline=now, status="Draft")
        for member in (self.ec1, self.ec2):
            TenderEvaluationCommitteeMember.objects.create(tender=self.tender, member=member, coi_attested=True,
                                                           coi_attested_at=now - timedelta(days=19))
        self.bid = TenderBid.objects.create(tender=self.tender, vendor_id=str(self.vendor.id), vendor_name="Sun Co", status="Awarded", bid_amount=90000)
        self.rival = TenderBid.objects.create(tender=self.tender, vendor_id="999", vendor_name="Moon Co", status="Not Awarded", bid_amount=95000)
        for bid, score in ((self.bid, 80), (self.rival, 60)):
            TenderBidEvaluation.objects.create(bid=bid, evaluator=self.ec1, stage="technical", submission_status="submitted", status="Scored",
                                               technical_score=score, total_score=score, justifications={"technical": "Meets the specification."},
                                               submitted_at=now - timedelta(days=15))
        # ec2 declared a conflict with Moon Co and still scored it; it never scored Sun Co.
        EvaluationConflictOfInterest.objects.create(evaluator=self.ec2, tender=self.tender, vendor_id="999", vendor_name="Moon Co")
        TenderBidEvaluation.objects.create(bid=self.rival, evaluator=self.ec2, stage="technical", submission_status="submitted", status="Scored",
                                           technical_score=50, total_score=50, justifications={}, submitted_at=now - timedelta(days=15))
        self.project = Project.objects.create(vendor_id=str(self.vendor.id), vendor_name="Sun Co", tech_type="SHS", region="Maseru",
                                              district="Maseru", project_reference="PRJ-P4-1", status="active", budget=90000,
                                              target_installations=10, tender=self.tender)
        self.install = InstallationReport.objects.create(project=self.project, vendor=self.vendor, gps_lat=-29.312345, gps_lng=27.481234,
                                                         serial_number="SN-P4-1", beneficiary_id="BEN-12345", beneficiary_name="Mpho Example",
                                                         beneficiary_phone="+26650000001", district="Maseru", status=InstallationStatus.VERIFIED,
                                                         household_type="female_headed")
        FieldVerification.objects.create(installation=self.install, field_officer=self.rmt, verification_status="Verified", beneficiary_gender="female",
                                         officer_latitude=-29.312345, officer_longitude=27.481234)
        self.claim = PaymentClaim.objects.create(project=self.project, vendor=self.vendor, claim_amount=9000, status="Completed", paid_at=now)

    def queue(self, user, report_type, fmt="csv", filters=None):
        self.client.force_authenticate(user)
        return self.client.post("/api/projects/reports/generate/", {"report_type": report_type, "format": fmt, "filters": filters or {}}, format="json")

    def run_report(self, user, report_type, fmt="csv", filters=None):
        queued = self.queue(user, report_type, fmt, filters)
        self.assertEqual(queued.status_code, status.HTTP_202_ACCEPTED, getattr(queued, "data", ""))
        report_queue.process()
        job = self.client.get(f"/api/projects/reports/{queued.data['id']}/").data
        self.assertEqual(job["status"], "ready", job.get("error"))
        response = self.client.get(job["downloadUrl"])
        return b"".join(response.streaming_content) if hasattr(response, "streaming_content") else response.content

    def rows(self, content):
        return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))

    def test_catalogue_access_by_role(self, *_):
        self.assertEqual(self.queue(self.vendor, "rmt_procurement_pack", "pdf").status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(self.queue(self.psc, "auditor_system_access").status_code, status.HTTP_403_FORBIDDEN)
        self.client.force_authenticate(self.ec1)
        ids = {t["id"] for t in self.client.get("/api/projects/reports/templates/").data}
        self.assertEqual(ids, {"ec_my_score_sheet"})
        tenders = [t["reference"] for t in self.client.get("/api/projects/reports/filter-options/").data["tenders"]]
        self.assertEqual(tenders, ["TND-P4-1"])

    def test_procurement_compliance_finds_conflicted_and_missing_scoring(self, *_):
        rows = self.rows(self.run_report(self.auditor, "auditor_procurement_compliance"))
        results = {(r["Tender"], r["Check"]): r["Result"] for r in rows}
        self.assertEqual(results[("TND-P4-1", "Committee of 3 to 5 members")], "Exception")
        self.assertEqual(results[("TND-P4-1", "Declared conflicts kept out of scoring")], "Exception")
        self.assertEqual(results[("TND-P4-1", "Every bid scored by every member")], "Exception")
        self.assertEqual(results[("TND-P4-1", "Technical scores justified")], "Exception")
        self.assertEqual(results[("TND-P4-1", "Standstill period kept")], "Exception")
        self.assertEqual(results[("TND-P4-1", "Publication approved by Super Admin")], "Exception")
        self.assertNotIn(("TND-P4-2", "Committee of 3 to 5 members"), results)

    def test_technical_evaluation_needs_a_tender_and_ranks_bids(self, *_):
        self.assertEqual(self.queue(self.tac, "tac_technical_evaluation", "excel").status_code, status.HTTP_400_BAD_REQUEST)
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(self.run_report(self.tac, "tac_technical_evaluation", "excel", {"tender": str(self.tender.id)})))
        ranking = [[c.value for c in row] for row in book["Ranking"].iter_rows(min_row=2)]
        self.assertEqual([(r[1], r[2], r[4]) for r in ranking], [(1, "Sun Co", 80), (2, "Moon Co", 55)])

    def test_score_sheet_shows_only_my_scores(self, *_):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(self.run_report(self.ec2, "ec_my_score_sheet", "excel")))
        scores = [[c.value for c in row] for row in book["My scores"].iter_rows(min_row=2)]
        self.assertEqual([(r[0], r[3]) for r in scores], [("TND-P4-1", "Moon Co")])

    def test_beneficiary_export_requires_a_reason_and_logs_it(self, *_):
        self.assertEqual(self.queue(self.rmt, "rmt_beneficiary_export").status_code, status.HTTP_400_BAD_REQUEST)
        rows = self.rows(self.run_report(self.rmt, "rmt_beneficiary_export", filters={"reason": "Grievance follow-up for household"}))
        self.assertEqual((rows[0]["Beneficiary ID"], rows[0]["Sex"]), ("BEN-12345", "Female"))
        logged = AuditLog.objects.filter(action="report_generated", details__report_type="rmt_beneficiary_export").first()
        self.assertEqual(logged.details["filters"]["reason"], "Grievance follow-up for household")

    def test_site_register_geojson_masks_coordinates_for_psc(self, *_):
        exact = json.loads(self.run_report(self.rmt, "site_register_gis", "geojson"))
        self.assertEqual(exact["features"][0]["geometry"]["coordinates"], [27.481234, -29.312345])
        masked = json.loads(self.run_report(self.psc, "site_register_gis", "geojson"))
        self.assertEqual(masked["features"][0]["geometry"]["coordinates"], [27.48, -29.31])
        self.assertEqual(masked["features"][0]["properties"]["Serial number"], "SN-P4-1")

    def test_iati_xml_lists_paid_claims_as_disbursements(self, *_):
        root = ET.fromstring(self.run_report(self.psc, "donor_iati_export", "xml"))
        transactions = root.findall("./iati-activity/transaction")
        self.assertEqual(len(transactions), 1)
        self.assertEqual(transactions[0].find("transaction-type").get("code"), "3")
        self.assertEqual(transactions[0].find("value").text, "9000.00")
        self.assertEqual(root.find("./iati-activity/recipient-country").get("code"), "LS")

    def test_project_audit_package_zip_with_checksums(self, *_):
        content = self.run_report(self.auditor, "auditor_project_package", "zip", {"project_id": str(self.project.id)})
        package = zipfile.ZipFile(io.BytesIO(content))
        names = set(package.namelist())
        self.assertTrue({"MANIFEST.txt", "workbook.xlsx", "lifecycle_dossier.pdf", "data/04_claims.csv"} <= names)
        manifest = package.read("MANIFEST.txt").decode()
        self.assertIn(f"{hashlib.sha256(package.read('data/04_claims.csv')).hexdigest()}  data/04_claims.csv", manifest)

    def test_system_access_counts_failed_sign_ins(self, *_):
        for _ in range(5):
            self.client.post("/api/users/auth/token/", {"username": "p4_rmt", "password": "wrong"}, format="json")
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(self.run_report(self.auditor, "auditor_system_access", "excel")))
        failed = [[c.value for c in row] for row in book["Failed sign-ins"].iter_rows(min_row=2)]
        self.assertEqual((failed[0][0], failed[0][1], failed[0][4]), ("p4_rmt", 5, "Review"))

    def test_vendor_scorecard_and_followups_run_for_their_audiences(self, *_):
        rows = self.rows(self.run_report(self.psc, "rmt_vendor_scorecard"))
        self.assertEqual((rows[0]["Vendor"], rows[0]["Verified"]), ("Sun Co", "1"))
        self.run_report(self.doe, "oversight_followup_tracker")
        self.run_report(self.vendor, "vendor_compliance_notices")

    def test_schedules_reject_reports_needing_a_tender_or_reason(self, *_):
        self.client.force_authenticate(self.admin)
        for report_type in ("tac_technical_evaluation", "rmt_beneficiary_export"):
            response = self.client.post("/api/projects/report-schedules/", {
                "name": "x", "report_type": report_type, "format": "excel", "frequency": "monthly", "run_day": 5,
                "prepared_by": self.rmt.id, "recipient_roles": ["Auditor"],
            }, format="json")
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST, report_type)
