"""First-wave reports: results framework, claims pipeline, audit chain and reconciliation, vendor and field reports."""

import csv
import io
from unittest import mock

from rest_framework import status
from rest_framework.test import APITestCase

from rbf.users.models import User
from . import report_queue
from .models import AuditLog, InstallationReport, InstallationStatus, Milestone, PaymentClaim, Project, ResultsIndicator, VerificationTask


def make_user(username, role, **extra):
    return User.objects.create_user(username=username, password="securePass123", role=role, status="Active", **extra)


def log(actor, action, claim):
    AuditLog.objects.create(actor=actor, actor_role=actor.role, action=action, entity_type="PaymentClaim", entity_id=str(claim.id))


@mock.patch("rbf.projects.report_pdf.render_report_pdf", return_value=b"%PDF-1.4 test")
class FirstWaveReportTests(APITestCase):
    def setUp(self):
        self.admin = make_user("p2_admin", "Platform Administrator (Super Admin)")
        self.rmt = make_user("p2_rmt", "RBF Management Team")
        self.tac = make_user("p2_tac", "TAC Member")
        self.psc = make_user("p2_psc", "Project Steering Committee")
        self.auditor = make_user("p2_auditor", "Auditor")
        self.fv = make_user("p2_fv", "Field Verifier", region="Maseru")
        self.vendor = make_user("p2_vendor", "Vendor", organization_name="Sun Co")
        self.other = make_user("p2_other", "Vendor", organization_name="Moon Co")
        self.project = Project.objects.create(
            vendor_id=str(self.vendor.id), vendor_name="Sun Co", tech_type="SHS", region="Maseru", district="Maseru",
            project_reference="PRJ-P2-1", status="active", budget=50000, target_installations=10,
        )
        self.other_project = Project.objects.create(
            vendor_id=str(self.other.id), vendor_name="Moon Co", tech_type="SHS", region="Berea", district="Berea",
            project_reference="PRJ-P2-2", status="active", budget=1000,
        )
        self.milestone = Milestone.objects.create(project=self.project, name="M1", milestone_number=1, percentage=20, amount=10000, amount_lsl=10000)
        self.clean = PaymentClaim.objects.create(project=self.project, vendor=self.vendor, milestone=self.milestone, claim_amount=10000, status="Completed")
        for actor, action in ((self.vendor, "payment_claim_submitted"), (self.rmt, "claim_rmt_approved"), (self.tac, "claim_tac_endorsed"),
                              (self.psc, "claim_psc_approved"), (self.admin, "finance_payment_processed"), (self.rmt, "payment_confirmed")):
            log(actor, action, self.clean)
        # RMT approves twice in the chain (segregation of duties) and the claim was paid with no PSC record.
        self.bad = PaymentClaim.objects.create(project=self.project, vendor=self.vendor, milestone=self.milestone, claim_amount=12000, status="Completed")
        log(self.vendor, "payment_claim_submitted", self.bad)
        log(self.rmt, "claim_rmt_approved", self.bad)
        log(self.rmt, "claim_tac_endorsed", self.bad)
        log(self.rmt, "payment_confirmed", self.bad)
        self.waiting = PaymentClaim.objects.create(project=self.project, vendor=self.vendor, milestone=self.milestone, claim_amount=10000, status="TAC Endorsed")
        self.other_claim = PaymentClaim.objects.create(project=self.other_project, vendor=self.other, claim_amount=500, status="Submitted")
        report = InstallationReport.objects.create(
            project=self.project, vendor=self.vendor, gps_lat=-29.31, gps_lng=27.48, serial_number="SN-P2-1", beneficiary_id="B-1",
            beneficiary_name="Mpho Example", beneficiary_phone="+26650000001", district="Maseru", status=InstallationStatus.SUBMITTED,
        )
        VerificationTask.objects.create(report=report, assigned_verifier=self.fv, vendor_lat=-29.31, vendor_lng=27.48)

    def run_report(self, user, report_type, fmt="csv", filters=None):
        self.client.force_authenticate(user)
        queued = self.client.post("/api/projects/reports/generate/", {"report_type": report_type, "format": fmt, "filters": filters or {}}, format="json")
        self.assertEqual(queued.status_code, status.HTTP_202_ACCEPTED, getattr(queued, "data", ""))
        report_queue.process()
        job = self.client.get(f"/api/projects/reports/{queued.data['id']}/").data
        self.assertEqual(job["status"], "ready", job.get("error"))
        response = self.client.get(job["downloadUrl"])
        return b"".join(response.streaming_content) if hasattr(response, "streaming_content") else response.content

    def rows(self, content):
        return list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))

    def sheet(self, content, prefix):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(content))
        ws = book[next(n for n in book.sheetnames if n.startswith(prefix))]
        header = [c.value for c in ws[1]]
        return [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]

    def test_payment_chain_flags_segregation_and_missing_psc(self, _pdf):
        rows = {int(r["Claim"]): r for r in self.rows(self.run_report(self.auditor, "auditor_payment_chain_audit"))}
        self.assertEqual(rows[self.clean.id]["Exceptions"], "No exception")
        self.assertEqual(rows[self.clean.id]["PSC approved by"], "p2_psc")
        bad = rows[self.bad.id]["Exceptions"]
        self.assertIn("Same person approved more than one step", bad)
        self.assertIn("Paid without a recorded PSC approval", bad)
        self.assertIn("No audit record for", bad)
        self.assertIn('TAC endorsed recorded by role "RBF Management Team"', bad)

    def test_reconciliation_reports_amount_differences(self, _pdf):
        rows = {int(r["Claim"]): r for r in self.rows(self.run_report(self.auditor, "auditor_reconciliation"))}
        self.assertIn("differs from milestone amount", rows[self.bad.id]["Result"])
        self.assertIn("no disbursement record", rows[self.clean.id]["Result"])

    def test_claims_pipeline_lists_only_open_claims(self, _pdf):
        rows = self.rows(self.run_report(self.rmt, "rmt_claims_pipeline"))
        self.assertEqual({int(r["Claim"]) for r in rows}, {self.waiting.id, self.other_claim.id})
        stage = {int(r["Claim"]): r["Waiting for"] for r in rows}
        self.assertEqual(stage[self.waiting.id], "Waiting for PSC")

    def test_results_report_uses_entered_targets(self, _pdf):
        ResultsIndicator.objects.create(code="OUT1", name="Households connected", unit="households", measure="amount_disbursed",
                                        baseline=0, target=100000)
        ResultsIndicator.objects.create(code="OUT9", name="Schools electrified", unit="schools", measure="manual", target=24, manual_actual=6)
        rows = {r["Code"]: r for r in self.sheet(self.run_report(self.psc, "programme_results", "excel"), "Results framework")}
        self.assertEqual(float(rows["OUT1"]["Actual to date"]), 22000.0)
        self.assertEqual(float(rows["OUT1"]["% of target"]), 22.0)
        self.assertEqual(float(rows["OUT9"]["% of target"]), 25.0)
        self.assertIn("Entered manually", rows["OUT9"]["Source"])

    def test_results_report_without_framework_lists_measures_and_says_so(self, _pdf):
        content = self.run_report(self.rmt, "programme_results", "excel")
        from openpyxl import load_workbook

        notes = " ".join(str(c.value) for c in load_workbook(io.BytesIO(content))["Summary"]["B"] if c.value)
        self.assertIn("No results framework has been entered yet", notes)

    def test_only_super_admin_edits_the_results_framework(self, _pdf):
        payload = {"code": "out2", "name": "Female-headed share", "measure": "female_headed_share", "target": 50}
        self.client.force_authenticate(self.rmt)
        self.assertEqual(self.client.post("/api/projects/results-indicators/", payload, format="json").status_code, status.HTTP_403_FORBIDDEN)
        self.client.force_authenticate(self.admin)
        created = self.client.post("/api/projects/results-indicators/", payload, format="json")
        self.assertEqual(created.status_code, status.HTTP_201_CREATED, created.data)
        self.assertEqual(created.data["code"], "OUT2")
        self.assertEqual(self.client.delete(f"/api/projects/results-indicators/{created.data['id']}/").status_code, status.HTTP_405_METHOD_NOT_ALLOWED)
        self.client.force_authenticate(self.psc)
        self.assertEqual(len(self.client.get("/api/projects/results-indicators/").data["results"]), 1)
        self.client.force_authenticate(self.vendor)
        self.assertEqual(self.client.get("/api/projects/results-indicators/").data["results"], [])

    def test_vendor_statement_shows_only_own_claims(self, _pdf):
        rows = self.rows(self.run_report(self.vendor, "vendor_claim_statement"))
        self.assertEqual({int(r["Claim"]) for r in rows}, {self.clean.id, self.bad.id, self.waiting.id})
        self.client.force_authenticate(self.vendor)
        forbidden = self.client.post("/api/projects/reports/generate/", {"report_type": "auditor_payment_chain_audit", "format": "csv"}, format="json")
        self.assertEqual(forbidden.status_code, status.HTTP_403_FORBIDDEN)

    def test_route_list_gives_verifier_full_contact_details(self, _pdf):
        rows = self.rows(self.run_report(self.fv, "fo_route_list"))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Beneficiary"], "Mpho Example")
        self.assertEqual(rows[0]["Phone"], "+26650000001")

    def test_psc_briefing_lists_decisions_needed(self, _pdf):
        rows = self.sheet(self.run_report(self.psc, "psc_briefing", "excel"), "Decisions needed")
        self.assertEqual([int(r["Claim"]) for r in rows], [self.waiting.id])
