"""Reports: access by role, data scope, filters, personal-data masking, background generation."""

import csv
import io
import time
from datetime import timedelta
from unittest import mock

from django.utils import timezone
from rest_framework import status
from rest_framework.test import APITestCase

from rbf.notifications.models import Notification
from rbf.users.models import User
from . import report_queue
from .models import (
    AuditLog,
    FieldVerification,
    GeneratedReport,
    InstallationReport,
    InstallationStatus,
    PaymentClaim,
    Project,
    VerificationTask,
)
from .report_catalogue import REGISTRY

ROLE_USERS = {
    "RBF Management Team": "rep_rmt",
    "Project Steering Committee": "rep_psc",
    "Auditor": "rep_auditor",
    "TAC Member": "rep_tac",
    "DoE Officer": "rep_doe",
    "Field Verifier": "rep_fv",
}


def make_user(username, role, **extra):
    return User.objects.create_user(username=username, password="securePass123", role=role, status="Active", **extra)


@mock.patch("rbf.projects.report_pdf.render_report_pdf", return_value=b"%PDF-1.4 test")
class ReportTests(APITestCase):
    def setUp(self):
        self.users = {role: make_user(name, role, region="Maseru") for role, name in ROLE_USERS.items()}
        self.users["DoE Officer"].districts = ["Maseru", "Berea"]
        self.users["DoE Officer"].save(update_fields=["districts"])
        self.admin = make_user("rep_admin", "Platform Administrator (Super Admin)")
        self.vendor = make_user("rep_vendor", "Vendor")
        self.other_vendor = make_user("rep_vendor2", "Vendor")
        self.maseru = Project.objects.create(
            vendor_id=str(self.vendor.id), vendor_name="Sun Co", tech_type="SHS", region="Maseru", district="Maseru",
            project_reference="PRJ-REP-1", status="active", budget=10000,
        )
        self.berea = Project.objects.create(
            vendor_id=str(self.vendor.id), vendor_name="Sun Co", tech_type="SHS", region="Berea", district="Berea",
            project_reference="PRJ-REP-2", status="active", budget=5000,
        )
        self.leribe = Project.objects.create(
            vendor_id=str(self.other_vendor.id), vendor_name="Moon Co", tech_type="GMG", region="Leribe", district="Leribe",
            project_reference="PRJ-REP-3", status="active", budget=20000,
        )
        self.paid = PaymentClaim.objects.create(project=self.maseru, vendor=self.vendor, claim_amount=1000, status="Completed")
        self.awaiting = PaymentClaim.objects.create(project=self.maseru, vendor=self.vendor, claim_amount=400, status="PSC Approved")
        self.legacy_paid = PaymentClaim.objects.create(project=self.leribe, vendor=self.other_vendor, claim_amount=250, status="Paid")
        old = PaymentClaim.objects.create(project=self.leribe, vendor=self.other_vendor, claim_amount=999, status="Completed")
        PaymentClaim.objects.filter(id=old.id).update(submitted_at=timezone.now() - timedelta(days=400))
        self.install = InstallationReport.objects.create(
            project=self.maseru, vendor=self.vendor, gps_lat=-29.312345, gps_lng=27.481234, serial_number="SN-REP-1",
            beneficiary_id="BEN-778899", district="Maseru", household_type="female_headed", status=InstallationStatus.VERIFIED,
        )
        VerificationTask.objects.create(report=self.install, vendor_lat=-29.3, vendor_lng=27.5)
        FieldVerification.objects.create(
            installation=self.install, field_officer=self.users["Field Verifier"], officer_latitude=-29.3,
            officer_longitude=27.5, location_match=False, verification_status="flagged", flag_reason="Panel missing",
        )

    # --- helpers --------------------------------------------------------------------

    def queue(self, user, report_type, fmt="csv", filters=None):
        self.client.force_authenticate(user)
        return self.client.post("/api/projects/reports/generate/", {"report_type": report_type, "format": fmt, "filters": filters or {}}, format="json")

    def generate(self, user, report_type, fmt="csv", filters=None):
        """Queue, let the worker run, then download. Returns the download response."""
        queued = self.queue(user, report_type, fmt, filters)
        if queued.status_code != status.HTTP_202_ACCEPTED:
            return queued
        report_queue.process()
        job = self.client.get(f"/api/projects/reports/{queued.data['id']}/").data
        self.assertEqual(job["status"], "ready", job.get("error"))
        return self.client.get(job["downloadUrl"])

    def content(self, response):
        return b"".join(response.streaming_content) if hasattr(response, "streaming_content") else response.content

    def csv_rows(self, response):
        return list(csv.reader(io.StringIO(self.content(response).decode("utf-8-sig"))))

    # --- catalogue and access -----------------------------------------------------------

    def test_every_registered_report_generates_in_every_format_for_its_roles(self, _pdf):
        from django.utils import timezone
        from rbf.tenders.models import Tender

        tender = Tender.objects.create(reference_number="TND-ALL-1", name="All reports", department="DoE", category="Energy",
                                       deadline=timezone.now(), status="Evaluation")
        for definition in REGISTRY.values():
            role = next(iter(definition.roles))
            filters = {"project_id": self.maseru.id} if definition.requires_project else {}
            if definition.requires_tender:
                filters["tender"] = str(tender.id)
            if definition.requires_reason:
                filters["reason"] = "Testing every report in the catalogue"
            user = self.users.get(role, self.admin)
            for fmt in definition.formats:
                response = self.generate(user, definition.id, fmt, filters)
                self.assertEqual(response.status_code, status.HTTP_200_OK, f"{definition.id} {fmt}")
                self.assertNotIn(b"placeholder", self.content(response).lower(), definition.id)

    def test_roles_cannot_generate_reports_outside_their_catalogue(self, _pdf):
        self.assertEqual(self.queue(self.vendor, "psc_vendor_payment_trail").status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(self.queue(self.users["DoE Officer"], "auditor_full_audit").status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(self.queue(self.users["Field Verifier"], "rmt_financial_disbursement").status_code, status.HTTP_403_FORBIDDEN)
        for removed in ("psc_quarterly_report", "psc_compliance_report"):
            self.assertEqual(self.queue(self.admin, removed).status_code, status.HTTP_400_BAD_REQUEST, removed)
        self.assertEqual(GeneratedReport.objects.count(), 0)

    def test_template_list_matches_role_and_declares_filters(self, _pdf):
        self.client.force_authenticate(self.users["DoE Officer"])
        templates = self.client.get("/api/projects/reports/templates/").data
        self.assertTrue(templates)
        ids = {t["id"] for t in templates}
        self.assertIn("doe_regional_progress", ids)
        self.assertFalse(ids & {"auditor_full_audit", "rmt_beneficiary_export", "rmt_procurement_pack"})
        self.assertTrue(all("filters" in t and t["method"] for t in templates))
        self.client.force_authenticate(self.vendor)
        vendor_ids = {t["id"] for t in self.client.get("/api/projects/reports/templates/").data}
        self.assertEqual(vendor_ids, {"rmt_kpi_project", "vendor_claim_statement", "vendor_installations", "vendor_compliance_notices",
                                      "rmt_energy_output_export"})

    # --- scope and filters --------------------------------------------------------------

    def test_doe_reports_cover_every_assigned_district_only(self, _pdf):
        rows = self.csv_rows(self.generate(self.users["DoE Officer"], "doe_regional_progress"))
        self.assertEqual({row[0] for row in rows[1:]}, {"PRJ-REP-1", "PRJ-REP-2"})

    def test_filter_options_follow_scope(self, _pdf):
        self.client.force_authenticate(self.users["DoE Officer"])
        options = self.client.get("/api/projects/reports/filter-options/").data
        self.assertEqual(options["districts"], ["Berea", "Maseru"])
        self.assertEqual([p["reference"] for p in options["projects"]], ["PRJ-REP-1", "PRJ-REP-2"])
        self.client.force_authenticate(self.other_vendor)
        self.assertEqual([p["reference"] for p in self.client.get("/api/projects/reports/filter-options/").data["projects"]], ["PRJ-REP-3"])

    def test_district_technology_and_vendor_filters_change_the_figures(self, _pdf):
        rmt = self.users["RBF Management Team"]
        everything = self.csv_rows(self.generate(rmt, "rmt_portfolio_summary"))
        self.assertEqual(len(everything) - 1, 3)
        berea = self.csv_rows(self.generate(rmt, "rmt_portfolio_summary", filters={"district": "Berea"}))
        self.assertEqual([r[0] for r in berea[1:]], ["PRJ-REP-2"])
        gmg = self.csv_rows(self.generate(rmt, "rmt_portfolio_summary", filters={"technology": "gmg"}))
        self.assertEqual([r[0] for r in gmg[1:]], ["PRJ-REP-3"])
        moon = self.csv_rows(self.generate(rmt, "rmt_financial_disbursement", filters={"vendor": str(self.other_vendor.id)}))
        self.assertEqual({r[2] for r in moon[1:]}, {"Moon Co"})

    def test_finance_totals_count_current_and_legacy_statuses(self, _pdf):
        rows = self.csv_rows(self.generate(self.users["RBF Management Team"], "rmt_financial_disbursement"))
        header, body = rows[0], rows[1:]
        stage = {row[0]: row[header.index("Stage")] for row in body}
        self.assertEqual(stage[str(self.paid.id)], "Paid")
        self.assertEqual(stage[str(self.legacy_paid.id)], "Paid")
        self.assertEqual(stage[str(self.awaiting.id)], "Approved, awaiting payment")

    def test_period_filter_is_applied_and_validated(self, _pdf):
        since = (timezone.localdate() - timedelta(days=30)).isoformat()
        rows = self.csv_rows(self.generate(self.users["RBF Management Team"], "rmt_financial_disbursement", filters={"from": since}))
        self.assertEqual(sorted(float(r[4]) for r in rows[1:]), [250.0, 400.0, 1000.0])
        bad = self.queue(self.users["RBF Management Team"], "rmt_financial_disbursement", filters={"from": "not-a-date"})
        self.assertEqual(bad.status_code, status.HTTP_400_BAD_REQUEST)

    def test_project_outside_scope_or_missing_is_refused_before_queueing(self, _pdf):
        self.assertEqual(self.queue(self.users["RBF Management Team"], "rmt_kpi_project").status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(
            self.queue(self.users["TAC Member"], "rmt_kpi_project", filters={"project_id": "999999"}).status_code,
            status.HTTP_403_FORBIDDEN,
        )

    # --- personal data ------------------------------------------------------------------

    def test_beneficiary_identity_and_gps_are_masked_for_tac_but_not_rmt(self, _pdf):
        from openpyxl import load_workbook

        def flagged_row(user):
            book = load_workbook(io.BytesIO(self.content(self.generate(user, "rmt_verification_project", "excel"))))
            sheet = book[next(name for name in book.sheetnames if name.startswith("Flagged and partial"))]
            header = [c.value for c in sheet[1]]
            return dict(zip(header, [c.value for c in sheet[2]])), [c.value for c in book["Summary"]["A"]]

        rmt_row, rmt_labels = flagged_row(self.users["RBF Management Team"])
        self.assertEqual(rmt_row["Beneficiary ID"], "BEN-778899")
        self.assertAlmostEqual(float(rmt_row["Latitude"]), -29.312345, places=5)
        self.assertNotIn("Data protection", rmt_labels)
        tac_row, tac_labels = flagged_row(self.users["TAC Member"])
        self.assertEqual(tac_row["Beneficiary ID"], "********99")
        self.assertEqual(float(tac_row["Latitude"]), -29.31)
        self.assertIn("Data protection", tac_labels)

    # --- background jobs, history, notifications --------------------------------------

    def test_jobs_queue_then_complete_and_notify(self, _pdf):
        queued = self.queue(self.users["RBF Management Team"], "rmt_portfolio_summary", "pdf")
        self.assertEqual(queued.status_code, status.HTTP_202_ACCEPTED)
        self.assertEqual(queued.data["status"], "queued")
        self.assertIsNone(queued.data["downloadUrl"])
        self.assertEqual(self.client.get(f"/api/projects/reports/{queued.data['id']}/download/").status_code, status.HTTP_409_CONFLICT)
        self.assertEqual(report_queue.process(), 1)
        job = self.client.get(f"/api/projects/reports/{queued.data['id']}/").data
        self.assertEqual(job["status"], "ready")
        self.assertTrue(Notification.objects.filter(event="report_ready", recipient_id=str(self.users["RBF Management Team"].id)).exists())
        self.assertTrue(AuditLog.objects.filter(action="report_generated").exists())

    def test_failed_job_is_recorded_and_notified(self, _pdf):
        queued = self.queue(self.users["RBF Management Team"], "rmt_portfolio_summary", "pdf")
        _pdf.side_effect = RuntimeError("renderer down")
        report_queue.process()
        job = self.client.get(f"/api/projects/reports/{queued.data['id']}/").data
        self.assertEqual(job["status"], "failed")
        self.assertNotIn("renderer down", job["error"])
        self.assertTrue(Notification.objects.filter(event="report_failed").exists())

    def test_stale_running_jobs_are_failed(self, _pdf):
        queued = self.queue(self.users["RBF Management Team"], "rmt_portfolio_summary")
        GeneratedReport.objects.filter(id=queued.data["id"]).update(status="running", started_at=timezone.now() - timedelta(hours=1))
        report_queue.process()
        self.assertEqual(GeneratedReport.objects.get(id=queued.data["id"]).status, "failed")

    def test_history_is_own_paged_and_searchable(self, _pdf):
        rmt = self.users["RBF Management Team"]
        for _ in range(27):
            self.queue(rmt, "rmt_portfolio_summary")
        self.queue(rmt, "rmt_anomaly_report")
        self.client.force_authenticate(rmt)
        first = self.client.get("/api/projects/reports/history/").data
        self.assertEqual((first["count"], len(first["results"])), (28, 25))
        self.assertEqual(len(self.client.get("/api/projects/reports/history/?page=2").data["results"]), 3)
        self.assertEqual(self.client.get("/api/projects/reports/history/?search=anomaly").data["count"], 1)
        self.client.force_authenticate(self.users["Project Steering Committee"])
        self.assertEqual(self.client.get("/api/projects/reports/history/").data["count"], 0)
        self.client.force_authenticate(self.users["Auditor"])
        self.assertEqual(self.client.get("/api/projects/reports/history/").data["count"], 28)

    def test_large_export_has_every_row(self, _pdf):
        AuditLog.objects.bulk_create([AuditLog(action=f"bulk_{i}", module="load") for i in range(10000)])
        started = time.monotonic()
        rows = self.csv_rows(self.generate(self.users["Auditor"], "auditor_full_audit"))
        self.assertGreaterEqual(len(rows) - 1, 10000)
        self.assertLess(time.monotonic() - started, 60)

    def test_excel_has_summary_and_table_sheets(self, _pdf):
        from openpyxl import load_workbook

        response = self.generate(self.users["RBF Management Team"], "rmt_financial_disbursement", "excel")
        workbook = load_workbook(io.BytesIO(self.content(response)))
        self.assertEqual(workbook.sheetnames, ["Summary", "By project", "Payment claims"])
        labels = [c.value for c in workbook["Summary"]["A"]]
        self.assertIn("Reference", labels)
        self.assertIn("Methodology", labels)
